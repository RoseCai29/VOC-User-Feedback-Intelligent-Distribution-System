"""
标注数据集加载器 —— 兼容真实数据集的字段差异。

支持 JSONL / JSON / CSV / TSV / XLSX，自动识别 text / category / sub_category / source
列名，也可由调用方显式指定；标签做轻量别名归一，减少人工整理数据集的成本。
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

TEXT_FIELDS = ("text", "raw_text", "content", "message", "feedback",
               "内容", "反馈内容", "用户反馈", "用户言论", "用户评论", "评论内容")
CATEGORY_FIELDS = ("category", "cat", "label", "一级分类", "主分类", "分类")
SUB_CATEGORY_FIELDS = ("sub_category", "sub_cat", "sub", "label_l2", "二级分类", "子分类")
SOURCE_FIELDS = ("source", "platform", "channel", "来源", "平台")

CATEGORY_ALIASES = {
    "建议": "作品建议", "用户建议": "作品建议", "需求": "作品建议",
    "吐槽": "舆情吐槽", "舆情": "舆情吐槽", "抱怨": "舆情吐槽",
    "售后": "售后报修", "报修": "售后报修", "维修": "售后报修",
    "好评": "正向反馈", "表扬": "正向反馈", "赞美": "正向反馈",
    "正面": "正向反馈", "positive": "正向反馈",
}

SUB_CATEGORY_ALIASES = {
    "功能需求": "功能建议", "需求建议": "功能建议",
    "UI": "界面优化", "界面": "界面优化", "体验": "界面优化",
    "推荐": "内容推荐", "算法推荐": "内容推荐",
    "质量": "产品质量", "品控": "产品质量",
    "客服": "服务态度", "服务": "服务态度", "态度": "服务态度",
    "价格": "价格问题", "定价": "价格问题",
    "硬件": "硬件故障", "故障": "硬件故障",
    "软件": "软件缺陷", "bug": "软件缺陷", "BUG": "软件缺陷", "崩溃": "软件缺陷",
    "物流": "物流问题", "发货": "物流问题", "快递": "物流问题",
    "商品好评": "产品好评", "产品称赞": "产品好评", "使用体验好": "产品好评",
    "售后好评": "服务好评", "服务称赞": "服务好评",
}


def read_records(path: str | Path) -> list[dict]:
    """按扩展名读取原始记录；表格文件返回每行一个 dict（键为表头）。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        with path.open(encoding="utf-8-sig") as f:
            return [json.loads(line) for line in f if line.strip()]
    if suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open(encoding="utf-8-sig", newline="") as f:
            return [dict(row) for row in csv.DictReader(f, delimiter=delimiter)]
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(rows)]
        records = [
            dict(zip(headers, row)) for row in rows
            if any(v is not None and str(v).strip() for v in row)
        ]
        wb.close()
        return records
    if suffix == ".xls":
        raise ValueError("不支持旧版 .xls，请另存为 .xlsx 或 .csv")
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    return data["examples"] if isinstance(data, dict) else data


def _pick(record: dict, explicit: str | None, candidates: tuple[str, ...]) -> str | None:
    if explicit:
        return explicit if explicit in record else None
    return next((k for k in candidates if k in record), None)


def _normalize_label(value, aliases: dict[str, str]) -> str:
    label = str(value or "").strip()
    return aliases.get(label, label)


def normalize_record(
    record: dict,
    text_field: str | None = None,
    category_field: str | None = None,
    sub_category_field: str | None = None,
    source_field: str | None = None,
) -> dict | None:
    """单条记录归一化为 {text, category, sub_category, source}，缺文本返回 None。"""
    tk = _pick(record, text_field, TEXT_FIELDS)
    if tk is None or not str(record.get(tk) or "").strip():
        return None
    ck = _pick(record, category_field, CATEGORY_FIELDS)
    sk = _pick(record, sub_category_field, SUB_CATEGORY_FIELDS)
    src_k = _pick(record, source_field, SOURCE_FIELDS)
    text = str(record.get(tk)).strip()
    normalized = {
        "text": text,
        "category": _normalize_label(record.get(ck), CATEGORY_ALIASES) if ck else "",
        "sub_category": _normalize_label(record.get(sk), SUB_CATEGORY_ALIASES) if sk else "",
        "source": str(record.get(src_k) or "社区").strip() if src_k else "社区",
    }
    # 可选的分组字段(如模板ID)：存在时透传，供按模板切分评估，避免同模板泄漏
    if record.get("template_id"):
        normalized["template_id"] = str(record["template_id"])
    return normalized


def load_examples(
    path: str | Path,
    text_field: str | None = None,
    category_field: str | None = None,
    sub_category_field: str | None = None,
    source_field: str | None = None,
) -> list[dict]:
    """加载并归一化数据集；字段名可用参数显式覆盖自动识别结果。"""
    examples = []
    for rec in read_records(path):
        normalized = normalize_record(
            rec, text_field, category_field, sub_category_field, source_field,
        )
        if normalized is not None:
            examples.append(normalized)
    return examples
